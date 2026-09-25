from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
import re
from typing import Any, Callable
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.dependencies import get_role_codes
from app.core.exceptions import BusinessException, bad_request, forbidden
from app.core.security import hash_password
from app.models.execution import ExecutionRecord
from app.models.import_job import ImportJob
from app.models.organization import Department, Organization
from app.models.project import Project, ProjectMember
from app.models.rbac import Role, UserRole
from app.models.schedule import ScheduleBooking
from app.models.task import Task, TaskAssignee
from app.models.user import User
from app.schemas.project import ProjectCreate
from app.schemas.task import TaskCreate
from app.schemas.user import UserCreate
from app.services.operation_log_service import log_operation
from app.services.notification_service import create_notification
from app.services.project_service import (
    PROJECT_CREATOR_ROLES,
    _add_initial_project_members,
    _generate_project_code,
    _validate_initial_project_members,
    get_department_manager,
    manageable_project_ids,
    visible_project_ids,
    _validate_project_manager,
)
from app.services.rbac_service import ensure_default_system_role
from app.services.task_service import _validate_assignees_with_parent, _validate_estimated_hours
from app.services.visibility_service import related_user_ids, has_global_project_access
from app.utils.time import beijing_now

RESOURCE_HEADERS = {
    "users": [
        ("employee_no", "员工号/登录账号*"),
        ("name", "姓名*"),
        ("password", "初始密码"),
        ("email", "邮箱"),
        ("department_code", "部门编码*"),
        ("organization_code", "组织编码"),
        ("supervisor_employee_no", "直属上级员工号*"),
        ("role_codes", "系统角色编码(逗号分隔)"),
        ("status", "状态"),
    ],
    "projects": [
        ("code", "项目编号（系统自动生成，导入值忽略）"),
        ("name", "项目名称*"),
        ("manager_employee_no", "项目经理员工号*"),
        ("member_employee_nos", "项目成员员工号*(逗号分隔)"),
        ("department_code", "部门编码*"),
        ("budget_hours", "项目总工时*"),
        ("planned_start", "计划开始*"),
        ("planned_end", "计划结束*"),
        ("description", "描述"),
        ("remark", "备注"),
    ],
    "tasks": [
        ("project_code", "项目编号*"),
        ("parent_task_name", "上级任务名称"),
        ("name", "任务名称*"),
        ("owner_employee_nos", "项目成员员工号*(逗号分隔)"),
        ("planned_start", "计划开始*"),
        ("planned_end", "计划结束*"),
        ("estimated_hours", "预计工时"),
        ("description", "描述"),
        ("remark", "备注"),
    ],
}

EXAMPLES = {
    "users": ["E10001", "张三", "", "zhangsan@example.com", "D001", "", "E10000", "project_member", "active"],
    "projects": ["P-2026-001", "示例项目", "E10001", "E10002,E10003", "D001", 160, date(2026, 10, 1), date(2026, 12, 31), "", ""],
    "tasks": ["P-2026-001", "", "需求分析", "E10002,E10003", date(2026, 10, 1), date(2026, 10, 3), 24, "", ""],
}


def _style_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="355F8D")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width


def _add_list_validation(book: Workbook, sheet, column: int, values: list[str]) -> None:
    lists = book["_lists"] if "_lists" in book.sheetnames else book.create_sheet("_lists")
    target_column = lists.max_column + 1 if lists.cell(1, 1).value else 1
    for row, value in enumerate(values, 1):
        lists.cell(row, target_column, value)
    column_letter = chr(64 + target_column)
    validation = DataValidation(
        type="list",
        formula1=f"'_lists'!${column_letter}$1:${column_letter}${len(values)}",
        allow_blank=True,
    )
    sheet.add_data_validation(validation)
    validation.add(f"{chr(64 + column)}2:{chr(64 + column)}5000")
    lists.sheet_state = "veryHidden"


def create_template(resource_type: str) -> bytes:
    headers = RESOURCE_HEADERS.get(resource_type)
    if not headers:
        raise bad_request("unsupported import resource")
    book = Workbook()
    sheet = book.active
    sheet.title = "导入数据"
    sheet.append([label for _, label in headers])
    sheet.append(EXAMPLES[resource_type])
    _style_sheet(sheet, [20] * len(headers))
    date_fields = {"planned_start", "planned_end"}
    for index, (field, _) in enumerate(headers, 1):
        if field in date_fields:
            sheet.cell(2, index).number_format = "yyyy-mm-dd"
    field_index = {field: index + 1 for index, (field, _) in enumerate(headers)}
    if "status" in field_index:
        status_values = {
            "users": ["active", "disabled"],
            "tasks": ["not_started", "running", "completed"],
        }[resource_type]
        _add_list_validation(book, sheet, field_index["status"], status_values)
    if "priority" in field_index:
        _add_list_validation(book, sheet, field_index["priority"], ["low", "medium", "high", "critical"])
    notes = book.create_sheet("填写说明", 1)
    notes.append(["规则", "说明"])
    notes.append(["必填字段", "标题包含 * 的列必须填写；请勿修改标题行。"])
    notes.append(["日期", "项目和任务计划日期请使用 Excel 日期单元格或 YYYY-MM-DD 格式。"])
    notes.append(["数字", "工时等数值请使用真实数字单元格，不要添加单位。"])
    if resource_type == "projects":
        notes.append(["项目成员", "创建项目时必须填写至少一名项目成员；多个员工号使用英文逗号分隔，项目经理无需重复填写。"])
    if resource_type == "tasks":
        notes.append(["任务项目成员", "至少填写一名当前有效项目成员；多个员工号使用英文逗号分隔。"])
    notes.append(["导入策略", "按行校验并导入；失败行会保留错误明细，成功行不会被回滚。"])
    _style_sheet(notes, [20, 76])
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def _clean(value: Any) -> Any:
    return value.strip() if isinstance(value, str) else value


def _to_datetime(value: Any, field: str) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, str):
        text = value.strip()
        for pattern in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
            try:
                return datetime.strptime(text, pattern)
            except ValueError:
                continue
    raise ValueError(f"{field} 必须是有效日期时间")


def _to_date(value: Any, field: str) -> date:
    return _to_datetime(value, field).date()


def _lookup(db: Session, model, column, value: Any, message: str):
    if value in (None, ""):
        return None
    statement = select(model).where(column == value)
    if model is User:
        statement = statement.where(User.is_deleted.is_(False))
    elif model is Project:
        statement = statement.where(Project.is_deleted.is_(False))
    elif model is Task:
        statement = statement.where(Task.is_deleted.is_(False))
    item = db.scalar(statement)
    if not item:
        raise ValueError(message)
    return item


def _required_lookup(db: Session, model, column, value: Any, message: str):
    if value in (None, ""):
        raise ValueError(message)
    item = _lookup(db, model, column, value, message)
    if not item:
        raise ValueError(message)
    return item


def _import_user(db: Session, row: dict[str, Any], operator: User) -> User:
    department = _required_lookup(
        db,
        Department,
        Department.code,
        row.get("department_code"),
        "部门编码不能为空且必须存在",
    )
    if department and department.status != "active":
        raise ValueError("不能把用户导入到已停用的部门")
    organization = None
    if row.get("organization_code"):
        if not department:
            raise ValueError("填写组织编码时必须同时填写部门编码")
        organization_filters = [Organization.code == row["organization_code"]]
        if department:
            organization_filters.append(Organization.department_id == department.id)
        organization = db.scalar(select(Organization).where(*organization_filters).limit(1))
        if not organization:
            raise ValueError("组织编码不存在或不属于所选部门")
        if organization.status != "active":
            raise ValueError("不能把用户导入到已停用的组织")
    supervisor = _required_lookup(
        db,
        User,
        User.employee_no,
        row.get("supervisor_employee_no"),
        "直属上级员工号必填且必须存在；最高级主管由数据库维护",
    )
    if supervisor and supervisor.status != "active":
        raise ValueError("直属上级已停用")
    role_codes = [item.strip() for item in str(row.get("role_codes") or "").split(",") if item.strip()]
    roles = db.scalars(select(Role).where(Role.code.in_(role_codes))).all() if role_codes else []
    missing_roles = set(role_codes) - {role.code for role in roles}
    if missing_roles:
        raise ValueError(f"角色编码不存在：{', '.join(sorted(missing_roles))}")
    if "super_admin" in role_codes:
        raise ValueError("超级管理员只能在系统初始化时配置，不能通过导入新增")
    payload = UserCreate(
        employee_no=row.get("employee_no"),
        name=row.get("name"),
        password=row.get("password") or settings.import_default_password,
        confirm_password=row.get("password") or settings.import_default_password,
        email=row.get("email") or None,
        department_id=department.id if department else None,
        organization_id=organization.id if organization else None,
        supervisor_id=supervisor.id if supervisor else None,
        status=row.get("status") or "active",
    )
    if payload.status not in {"active", "disabled"}:
        raise ValueError("用户状态必须是 active 或 disabled")
    if db.scalar(select(User.id).where(User.employee_no == payload.employee_no)):
        raise ValueError("员工号已存在")
    values = payload.model_dump(
        exclude={
            "employee_no",
            "password",
            "confirm_password",
            "role_ids",
        }
    )
    item = User(
        **values,
        employee_no=payload.employee_no,
        password_hash=hash_password(payload.password),
    )
    db.add(item)
    db.flush()
    db.add_all(
        [
            UserRole(user_id=item.id, role_id=role.id)
            for role in roles
        ]
    )
    ensure_default_system_role(db, item.id)
    return item


def _import_project(db: Session, row: dict[str, Any], operator: User) -> Project:
    manager = _required_lookup(db, User, User.employee_no, row.get("manager_employee_no"), "项目经理员工号不能为空且必须存在")
    department = _required_lookup(db, Department, Department.code, row.get("department_code"), "部门编码不能为空且必须存在")
    operator_roles = get_role_codes(db, operator.id)
    if not has_global_project_access(db, operator) and (manager.id != operator.id or not (operator_roles & PROJECT_CREATOR_ROLES)):
        raise ValueError("项目只能由项目经理本人导入")
    _validate_project_manager(db, manager.id, department.id)
    if manager.department_id != department.id:
        raise ValueError("项目经理必须属于项目所属部门")
    if department.status != "active":
        raise ValueError("不能向已停用的部门导入项目")
    member_employee_nos = [
        value.strip()
        for value in re.split(r"[,，;；]", str(row.get("member_employee_nos") or ""))
        if value.strip()
    ]
    member_employee_nos = list(dict.fromkeys(member_employee_nos))
    if not member_employee_nos:
        raise ValueError("创建项目时必须指定至少一名项目成员")
    if manager.employee_no in member_employee_nos:
        raise ValueError("项目经理无需在项目成员员工号中重复填写")
    member_users = db.scalars(
        select(User).where(User.employee_no.in_(member_employee_nos))
    ).all()
    member_users_by_no = {member.employee_no: member for member in member_users}
    missing_member_nos = [
        employee_no
        for employee_no in member_employee_nos
        if employee_no not in member_users_by_no
    ]
    if missing_member_nos:
        raise ValueError(f"项目成员员工号不存在：{', '.join(missing_member_nos)}")
    ordered_member_users = [
        member_users_by_no[employee_no] for employee_no in member_employee_nos
    ]
    try:
        _validate_initial_project_members(
            db, [member.id for member in ordered_member_users]
        )
    except BusinessException as exc:
        raise ValueError(exc.message) from exc
    payload = ProjectCreate(
        name=row.get("name"),
        manager_id=manager.id,
        member_ids=[member.id for member in ordered_member_users],
        department_id=department.id,
        budget_hours=Decimal(str(row.get("budget_hours") or 0)),
        planned_start=_to_date(row.get("planned_start"), "计划开始"),
        planned_end=_to_date(row.get("planned_end"), "计划结束"),
        description=row.get("description") or None,
        remark=row.get("remark") or None,
    )
    try:
        approver = get_department_manager(db, department.id)
    except BusinessException as exc:
        raise ValueError(exc.message) from exc
    item = Project(
        **payload.model_dump(exclude={"member_ids"}),
        code=_generate_project_code(db),
        status="not_started",
        approval_status="pending",
        created_by=operator.id,
        approver_id=approver.id,
    )
    db.add(item)
    db.flush()
    _add_initial_project_members(db, item, manager, ordered_member_users)
    db.flush()
    create_notification(
        db,
        approver.id,
        "project_approval_required",
        "导入项目等待部门主管审批",
        f"{operator.name} 导入了项目 {item.code} - {item.name}，"
        f"请审核项目与 {item.budget_hours} 小时工时额度。",
        level="warning",
        related_type="project",
        related_id=item.id,
    )
    return item


def _import_task(db: Session, row: dict[str, Any], operator: User) -> Task:
    project = _required_lookup(db, Project, Project.code, row.get("project_code"), "项目编号不能为空且必须存在")
    operator_roles = get_role_codes(db, operator.id)
    if not has_global_project_access(db, operator) and project.manager_id != operator.id:
        raise ValueError("任务只能由该项目的项目经理本人导入")
    if project.approval_status != "approved":
        raise ValueError("项目尚未通过审批，不能导入任务")
    if project.status == "completed":
        raise ValueError("已完成项目不能导入任务")
    owner_employee_nos = [
        value.strip()
        for value in re.split(r"[,，;；]", str(row.get("owner_employee_nos") or ""))
        if value.strip()
    ]
    owner_employee_nos = list(dict.fromkeys(owner_employee_nos))
    if not owner_employee_nos:
        raise ValueError("至少填写一名任务项目成员的员工号")
    owners = db.scalars(
        select(User).where(
            User.employee_no.in_(owner_employee_nos),
            User.status == "active",
            User.is_deleted.is_(False),
        )
    ).all()
    owners_by_no = {owner.employee_no: owner for owner in owners}
    missing_owner_nos = [
        employee_no
        for employee_no in owner_employee_nos
        if employee_no not in owners_by_no
    ]
    if missing_owner_nos:
        raise ValueError(
            f"任务项目成员员工号不存在或已停用：{', '.join(missing_owner_nos)}"
        )
    ordered_owners = [owners_by_no[employee_no] for employee_no in owner_employee_nos]
    active_member_ids = set(
        db.scalars(
            select(ProjectMember.user_id).where(
                ProjectMember.project_id == project.id,
                ProjectMember.left_at.is_(None),
            )
        ).all()
    )
    invalid_owners = [
        owner.employee_no for owner in ordered_owners if owner.id not in active_member_ids
    ]
    if invalid_owners:
        raise ValueError(
            f"任务项目成员必须是当前有效项目成员：{', '.join(invalid_owners)}"
        )
    parent = None
    if row.get("parent_task_name"):
        parent_matches = list(
            db.scalars(
                select(Task).where(
                    Task.project_id == project.id,
                    Task.name == row["parent_task_name"],
                    Task.is_deleted.is_(False),
                )
            ).all()
        )
        if not parent_matches:
            raise ValueError("上级任务名称不存在")
        if len(parent_matches) > 1:
            raise ValueError("上级任务名称不唯一，请先调整任务名称后再导入")
        parent = parent_matches[0]
        if parent.status == "completed":
            raise ValueError("不能在已完成的任务下导入子任务")
        if db.scalar(
            select(ExecutionRecord.id).where(
                ExecutionRecord.task_id == parent.id,
                ExecutionRecord.is_deleted.is_(False),
            ).limit(1)
        ):
            raise ValueError("已有执行记录的任务不能再作为汇总任务")
        if db.scalar(
            select(ScheduleBooking.id).where(
                ScheduleBooking.task_id == parent.id,
                ScheduleBooking.status.notin_({"rejected", "cancelled", "withdrawn"}),
            ).limit(1)
        ):
            raise ValueError("已有有效预约记录的任务不能再作为汇总任务")
    payload = TaskCreate(
        project_id=project.id,
        parent_id=parent.id if parent else None,
        name=row.get("name"),
        owner_ids=[owner.id for owner in ordered_owners],
        planned_start=_to_date(row.get("planned_start"), "计划开始"),
        planned_end=_to_date(row.get("planned_end"), "计划结束"),
        estimated_hours=Decimal(str(row.get("estimated_hours") or 0)),
        description=row.get("description") or None,
        remark=row.get("remark") or None,
    )
    if (
        payload.planned_start < project.planned_start
        or payload.planned_end > project.planned_end
    ):
        raise ValueError("任务计划时间必须位于项目计划日期范围内")
    if parent and (
        payload.planned_start < parent.planned_start
        or payload.planned_end > parent.planned_end
    ):
        raise ValueError("子任务计划时间必须位于父任务计划日期范围内")
    _validate_assignees_with_parent(
        db, parent, [owner.id for owner in ordered_owners]
    )
    _validate_estimated_hours(
        db,
        project.id,
        parent.id if parent else None,
        payload.estimated_hours,
    )
    item = Task(
        **payload.model_dump(exclude={"owner_ids"}),
        owner_id=ordered_owners[0].id,
        priority="medium",
        status="not_started",
    )
    db.add(item)
    db.flush()
    for owner in ordered_owners:
        db.add(TaskAssignee(task_id=item.id, user_id=owner.id))
        if owner.id != operator.id:
            create_notification(
                db,
                owner.id,
                "task_assigned",
                "你收到了一项新任务",
                f"任务“{item.name}”已通过导入分配给你。",
                related_type="task",
                related_id=item.id,
            )
    return item


IMPORT_HANDLERS: dict[str, Callable[[Session, dict[str, Any], User], Any]] = {
    "users": _import_user,
    "projects": _import_project,
    "tasks": _import_task,
}


def _assert_import_scope(
    db: Session, resource_type: str, row: dict[str, Any], operator: User
) -> None:
    roles = get_role_codes(db, operator.id)
    if roles & {"super_admin", "department_manager"}:
        return
    if resource_type in {"users", "projects"}:
        department_code = row.get("department_code")
        department_id = db.scalar(
            select(Department.id).where(Department.code == department_code)
        ) if department_code else None
        if not operator.department_id or department_id != operator.department_id:
            raise ValueError("只能导入到当前操作人所属部门")
    if resource_type == "tasks":
        project_id = db.scalar(select(Project.id).where(Project.code == row.get("project_code")))
        manageable = manageable_project_ids(db, operator)
        if not project_id or manageable is not None and project_id not in manageable:
            raise ValueError("任务所属项目不在当前操作人的数据范围内")


def import_workbook(
    db: Session, resource_type: str, filename: str, content: bytes, operator: User
) -> ImportJob:
    headers = RESOURCE_HEADERS.get(resource_type)
    handler = IMPORT_HANDLERS.get(resource_type)
    if not headers or not handler:
        raise bad_request("unsupported import resource")
    if resource_type == "users" and "super_admin" not in get_role_codes(db, operator.id):
        raise forbidden("只有超级管理员可以导入用户和分配系统角色")
    if not filename.lower().endswith(".xlsx"):
        raise bad_request("only .xlsx files are supported")
    if len(content) > settings.import_max_mb * 1024 * 1024:
        raise bad_request(f"file exceeds {settings.import_max_mb} MB")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (OSError, ValueError, KeyError, BadZipFile, InvalidFileException) as exc:
        raise bad_request("invalid Excel workbook") from exc
    if "导入数据" not in workbook.sheetnames:
        raise bad_request("workbook is missing sheet: 导入数据")
    sheet = workbook["导入数据"]
    expected_labels = [label for _, label in headers]
    actual_labels = [cell.value for cell in sheet[1]][: len(expected_labels)]
    if actual_labels != expected_labels:
        raise bad_request("Excel headers do not match the selected template")
    job = ImportJob(
        resource_type=resource_type,
        original_filename=filename[:255],
        operator_id=operator.id,
        status="processing",
    )
    db.add(job)
    db.flush()
    errors: list[dict] = []
    success_rows = 0
    failed_rows = 0
    total_rows = 0
    field_names = [field for field, _ in headers]
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if not any(value not in (None, "") for value in values):
            continue
        total_rows += 1
        row = {field: _clean(values[index]) if index < len(values) else None for index, field in enumerate(field_names)}
        try:
            with db.begin_nested():
                _assert_import_scope(db, resource_type, row, operator)
                handler(db, row, operator)
            success_rows += 1
        except ValidationError as exc:
            failed_rows += 1
            first = exc.errors()[0]
            if len(errors) < 500:
                errors.append({"row": row_number, "field": ".".join(map(str, first["loc"])), "message": first["msg"]})
        except BusinessException as exc:
            failed_rows += 1
            if len(errors) < 500:
                errors.append({"row": row_number, "field": None, "message": exc.message})
        except (ValueError, TypeError, IntegrityError) as exc:
            failed_rows += 1
            if len(errors) < 500:
                errors.append({"row": row_number, "field": None, "message": str(exc)[:500]})
    if failed_rows > 500:
        errors[-1] = {
            "row": 0,
            "field": None,
            "message": f"另有 {failed_rows - 499} 条错误未展示；导入仍已处理全部数据行",
        }
    job.total_rows = total_rows
    job.success_rows = success_rows
    job.failed_rows = failed_rows
    job.errors = errors
    job.status = "completed" if not failed_rows else ("partial" if success_rows else "failed")
    log_operation(
        db,
        operator_id=operator.id,
        module="data_exchange",
        action="import",
        object_type="import_job",
        object_id=job.id,
        after_data={
            "resource_type": resource_type,
            "filename": filename,
            "total_rows": total_rows,
            "success_rows": success_rows,
            "failed_rows": failed_rows,
        },
    )
    db.commit()
    db.refresh(job)
    return job


def list_import_jobs(db: Session, page: int, page_size: int, resource_type: str | None) -> dict:
    filters = [ImportJob.resource_type == resource_type] if resource_type else []
    total = db.scalar(select(func.count(ImportJob.id)).where(*filters)) or 0
    rows = db.execute(
        select(ImportJob, User.name.label("operator_name"))
        .join(User, User.id == ImportJob.operator_id)
        .where(*filters)
        .order_by(ImportJob.created_at.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    items = []
    for item, operator_name in rows:
        data = {column.name: getattr(item, column.name) for column in ImportJob.__table__.columns}
        data["operator_name"] = operator_name
        items.append(data)
    return {"items": items, "total": total, "page": page, "page_size": page_size}


def _export_book(
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    date_columns: set[int],
    datetime_columns: set[int] | None = None,
) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = title[:31]
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    _style_sheet(sheet, [18] * len(headers))
    for row in sheet.iter_rows(min_row=2):
        for index in date_columns:
            row[index - 1].number_format = "yyyy-mm-dd"
        for index in datetime_columns or set():
            row[index - 1].number_format = "yyyy-mm-dd hh:mm"
    summary = book.create_sheet("导出说明")
    summary.append(["项目", "内容"])
    summary.append(["报表", title])
    summary.append(["生成时间", beijing_now()])
    summary["B3"].number_format = "yyyy-mm-dd hh:mm:ss"
    summary.append(["数据行数", len(rows)])
    _style_sheet(summary, [18, 32])
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def _validate_export_range(start_date: date, end_date: date) -> None:
    if end_date < start_date:
        raise bad_request("end_date must be on or after start_date")
    if (end_date - start_date).days > 366:
        raise bad_request("export range cannot exceed 366 days")


def _export_filters(db: Session, user: User, project_column, user_column, *, project_id: int | None = None, personnel_keyword: str | None = None, organization_keyword: str | None = None) -> list:
    """Export related projects OR own/subordinate records, never other teams' unrelated work."""
    projects = visible_project_ids(db, user)
    people = related_user_ids(db, user)
    filters = []
    if projects is not None and people is not None:
        filters.append(or_(project_column.in_(projects or {-1}), user_column.in_(people or {-1})))
    if project_id:
        filters.append(project_column == project_id)
    if personnel_keyword and personnel_keyword.strip():
        term = f"%{personnel_keyword.strip()}%"
        filters.append(or_(User.name.like(term), User.employee_no.like(term)))
    if organization_keyword and organization_keyword.strip():
        term = f"%{organization_keyword.strip()}%"
        filters.append(or_(
            User.department_id.in_(select(Department.id).where(Department.name.like(term))),
            User.organization_id.in_(select(Organization.id).where(Organization.name.like(term))),
        ))
    return filters


def export_schedules(db: Session, user: User, start_date: date, end_date: date, **query) -> bytes:
    _validate_export_range(start_date, end_date)
    filters = [
        Project.is_deleted.is_(False),
        Task.is_deleted.is_(False),
        ScheduleBooking.end_time > datetime.combine(start_date, time.min),
        ScheduleBooking.start_time < datetime.combine(end_date + timedelta(days=1), time.min),
    ]
    filters.extend(_export_filters(db, user, ScheduleBooking.project_id, ScheduleBooking.user_id, **query))
    rows = db.execute(
        select(ScheduleBooking, User.name, Project.code, Project.name, Task.name)
        .join(User, User.id == ScheduleBooking.user_id)
        .join(Project, Project.id == ScheduleBooking.project_id)
        .join(Task, Task.id == ScheduleBooking.task_id)
        .where(*filters)
        .order_by(ScheduleBooking.start_time, User.employee_no.asc(), User.id.asc())
    ).all()
    values = [
        [item.id, user_name, code, project_name, task_name, item.start_time, item.end_time, float(item.planned_hours), item.status, item.remark]
        for item, user_name, code, project_name, task_name in rows
    ]
    return _export_book(
        "排期明细",
        ["预约ID", "人员", "项目编号", "项目", "任务", "开始", "结束", "计划工时", "状态", "备注"],
        values,
        set(),
        {6, 7},
    )


def export_executions(db: Session, user: User, start_date: date, end_date: date, **query) -> bytes:
    _validate_export_range(start_date, end_date)
    filters = [
        ExecutionRecord.is_deleted.is_(False),
        Task.is_deleted.is_(False),
        Project.is_deleted.is_(False),
        func.coalesce(
            ExecutionRecord.actual_end,
            ExecutionRecord.actual_start,
        )
        >= start_date,
        ExecutionRecord.actual_start <= end_date,
    ]
    filters.extend(_export_filters(db, user, Task.project_id, ExecutionRecord.user_id, **query))
    rows = db.execute(
        select(ExecutionRecord, User.name, Project.code, Project.name, Task.name)
        .join(User, User.id == ExecutionRecord.user_id)
        .join(Task, Task.id == ExecutionRecord.task_id)
        .join(Project, Project.id == Task.project_id)
        .where(*filters)
        .order_by(ExecutionRecord.actual_start)
    ).all()
    values = [
        [item.id, user_name, code, project_name, task_name, item.actual_start, item.actual_end, float(item.actual_hours), item.status, item.description]
        for item, user_name, code, project_name, task_name in rows
    ]
    return _export_book("执行明细", ["执行ID", "人员", "项目编号", "项目", "任务", "实际开始", "实际结束", "实际工时", "状态", "执行说明"], values, {6, 7})


def export_process_report(db: Session, user: User, start_date: date, end_date: date, **query) -> bytes:
    _validate_export_range(start_date, end_date)
    from app.services.report_service import process_report

    report = process_report(
        db, user, 1, 100000, start_date=start_date, end_date=end_date, **query
    )

    def period_text(start, end, *, open_ended: bool = False) -> str:
        if not start:
            return "—"
        return f"{start} 至 {end if end else ('进行中' if open_ended else '—')}"

    values = [
        [
            item["project_name"], item.get("task_path"), item["owner_name"],
            period_text(item["planned_start"], item["planned_end"]),
            period_text(item.get("actual_start"), item.get("actual_end"), open_ended=True),
            float(item["estimated_hours"] or 0), float(item["actual_hours"] or 0),
            float(item["achievement_rate"]) if item.get("achievement_rate") is not None else None,
            float(item["achievement_quality"]) if item.get("achievement_quality") is not None else None,
            item["effective_status"],
        ]
        for item in report["items"]
    ]
    return _export_book("项目过程报表", ["项目", "任务层级", "负责人", "计划工期", "实际工期", "预计工时", "实际工时", "达成率", "达成质量", "状态"], values, set())
